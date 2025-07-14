from docx import Document
import os

def run(nombre_archivo: str, contenido: str, workspace: str):
    doc = Document()
    doc.add_heading('Resumen Generado', level=1)
    for linea in contenido.strip().split("\n"):
        doc.add_paragraph(linea)

    path_output = f"storage/workspaces/{workspace}/output"
    os.makedirs(path_output, exist_ok=True)
    file_path = os.path.join(path_output, f"{nombre_archivo}.docx")
    doc.save(file_path)
    return file_path
