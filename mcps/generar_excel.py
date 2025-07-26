import os
import openpyxl
from openpyxl.utils import get_column_letter

def run(nombre_archivo, tabla, workspace):
    output_dir = os.path.join("storage", "workspaces", workspace, "output")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f"{nombre_archivo}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja 1"

    if isinstance(tabla, list):
        # Si es lista de diccionarios
        if all(isinstance(fila, dict) for fila in tabla):
            headers = list(tabla[0].keys())
            ws.append(headers)
            for fila in tabla:
                ws.append([fila.get(col, "") for col in headers])

        # Si es lista de listas
        elif all(isinstance(fila, list) for fila in tabla):
            for fila in tabla:
                ws.append(fila)

        else:
            ws["A1"] = "⚠️ Formato de tabla no compatible."
    else:
        ws["A1"] = "⚠️ La tabla no tiene formato de lista."

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(15, max_len + 2)

    wb.save(output_path)
    return output_path
