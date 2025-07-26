import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import os


def run(workspace: str, prompt: str = "") -> str:
    archivo_path = f"storage/workspaces/{workspace}/documents"
    excel_files = [f for f in os.listdir(archivo_path) if f.endswith((".xlsx", ".xlsm", ".xls"))]

    if not excel_files:
        raise FileNotFoundError("No se encontró ningún archivo Excel en el workspace.")

    ruta_excel = os.path.join(archivo_path, excel_files[0])
    df = pd.read_excel(ruta_excel)

    # Columnas posibles
    columnas_objetivo = []
    prompt_lower = prompt.lower()

    if "duplicadas" in prompt_lower:
        columnas_objetivo.append("Duplicado")
    if "apócrifas" in prompt_lower or "apocrifas" in prompt_lower:
        columnas_objetivo.append("Apócrifa")

    # Si no se detecta nada, por defecto resaltar duplicadas
    if not columnas_objetivo:
        columnas_objetivo = ["Duplicado"]

    wb = load_workbook(ruta_excel)
    ws = wb.active

    amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    for i, row in enumerate(df.itertuples(index=False), start=2):
        for col in columnas_objetivo:
            try:
                if getattr(row, col) == "Sí":
                    col_index = df.columns.get_loc(col) + 1
                    ws.cell(row=i, column=col_index).fill = amarillo
            except AttributeError:
                continue

    salida = f"storage/workspaces/{workspace}/output/resaltado_facturas.xlsx"
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    wb.save(salida)
    return salida
