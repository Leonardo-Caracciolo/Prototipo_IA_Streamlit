from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import os


def run(workspace: str, historial: list) -> str:
    doc = Document()

    # Título
    titulo = doc.add_heading(f"Resumen de Conversación - Workspace '{workspace}'", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Fecha
    fecha = doc.add_paragraph(datetime.now().strftime("%d/%m/%Y %H:%M"))
    fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph("\n")

    # Historial
    for pregunta, respuesta in historial:
        p_pregunta = doc.add_paragraph()
        p_pregunta.add_run("🧑 Usuario: ").bold = True
        p_pregunta.add_run(pregunta)

        p_respuesta = doc.add_paragraph()
        p_respuesta.add_run("🤖 TaxMiner: ").bold = True
        p_respuesta.add_run(respuesta)

        doc.add_paragraph("—" * 30)

    # Guardar archivo
    carpeta_salida = f"storage/workspaces/{workspace}/output"
    os.makedirs(carpeta_salida, exist_ok=True)

    nombre_archivo = f"{carpeta_salida}/resumen_conversacion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(nombre_archivo)
    return nombre_archivo
